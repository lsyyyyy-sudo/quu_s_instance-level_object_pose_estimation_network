"""生成极简版汇报 PPT（按 docs/HUMAN_LOG.md 末尾「交付」的三条主线组织）。

骨架完全照人类在 HUMAN_LOG 最后写的结构：
    （1）数据集主线  （2）模型主线  （3）效果
另加封面、任务定义、已知缺陷、教训与下一步，共 10 页。

⚠️ 中文字体的坑：python-pptx 只设 ``run.font.name`` 只影响**西文**，
东亚文字走 ``<a:ea typeface="...">``。必须两个都设，否则中文会退化成默认字体（甚至方框）。

用法::

    python scripts/make_ppt_minimal.py                  # 默认输出 docs/ppt/minimal.pptx
    python scripts/make_ppt_minimal.py --out x.pptx
"""

from __future__ import annotations

import argparse
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
FIG = os.path.join(ROOT, "docs", "figures")
RES = os.path.join(ROOT, "data", "results")

CJK = "Microsoft YaHei"
INK = RGBColor(0x1A, 0x1A, 0x1E)
MUTED = RGBColor(0x60, 0x60, 0x6C)
ACCENT = RGBColor(0x1F, 0x6F, 0xEB)
WARN = RGBColor(0xC0, 0x39, 0x2B)
GOOD = RGBColor(0x1E, 0x7A, 0x3C)

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)


def _font(run, size, bold=False, color=INK, name=CJK):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def bar(slide, y=Inches(0.0), h=Inches(0.10), color=ACCENT):
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), y, SLIDE_W, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def title(slide, text, sub=None):
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.30), SLIDE_W - Inches(1.2), Inches(0.75))
    tf = tb.text_frame
    tf.word_wrap = True
    _font(tf.paragraphs[0].add_run(), 28, bold=True)
    tf.paragraphs[0].runs[0].text = text
    if sub:
        tb2 = slide.shapes.add_textbox(Inches(0.6), Inches(1.00), SLIDE_W - Inches(1.2), Inches(0.40))
        tf2 = tb2.text_frame
        tf2.word_wrap = True
        _font(tf2.paragraphs[0].add_run(), 13, color=MUTED)
        tf2.paragraphs[0].runs[0].text = sub


def bullets(slide, items, left, top, width, height, size=15, space=8):
    """items: [(文本, 颜色 or None, 是否加粗)]"""
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for it in items:
        text = it[0] if isinstance(it, tuple) else it
        col = it[1] if isinstance(it, tuple) and len(it) > 1 and it[1] else INK
        bold = it[2] if isinstance(it, tuple) and len(it) > 2 else False
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(space)
        _font(p.add_run(), size, bold=bold, color=col)
        p.runs[0].text = text
    return tb


def picture_fit(slide, path, box, caption=None, cap_size=11):
    """把图等比缩放进 box=(l,t,w,h)（EMU），居中。"""
    if not os.path.isfile(path):
        print(f"  ⚠️ 图不存在，跳过: {path}")
        return None
    with Image.open(path) as im:
        iw, ih = im.size
    l, t, w, h = box
    scale = min(w / iw, h / ih)
    nw, nh = int(iw * scale), int(ih * scale)
    left = int(l + (w - nw) / 2)
    top = int(t + (h - nh) / 2)
    pic = slide.shapes.add_picture(path, Emu(left), Emu(top), Emu(nw), Emu(nh))
    if caption:
        cap = slide.shapes.add_textbox(Emu(l), Emu(t + h + Emu(Inches(0.04))), Emu(w), Inches(0.5))
        tf = cap.text_frame
        tf.word_wrap = True
        _font(tf.paragraphs[0].add_run(), cap_size, color=MUTED)
        tf.paragraphs[0].runs[0].text = caption
    return pic


def crop_region(src, dst, frac):
    """按相对比例裁出图的某一块（PPT 排版用：把竖排的多帧长条拆成单帧）。

    frac = (left, top, right, bottom)，都是 0~1 的相对比例。
    """
    os.makedirs(os.path.dirname(dst), exist_ok=True)
    with Image.open(src) as im:
        w, h = im.size
        l, t, r, b = frac
        im.crop((int(w * l), int(h * t), int(w * r), int(h * b))).save(dst, quality=95)
    return dst


def table(slide, rows, left, top, width, height, col_w=None, size=12.5, head_size=13):
    nrow, ncol = len(rows), len(rows[0])
    shp = slide.shapes.add_table(nrow, ncol, left, top, width, height)
    tbl = shp.table
    if col_w:
        total = sum(col_w)
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Emu(int(width * cw / total))
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            tf = cell.text_frame
            tf.word_wrap = True
            _font(tf.paragraphs[0].add_run(),
                  head_size if r == 0 else size,
                  bold=(r == 0),
                  color=RGBColor(0xFF, 0xFF, 0xFF) if r == 0 else INK)
            tf.paragraphs[0].runs[0].text = str(val)
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0x2B, 0x2B, 0x33)
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xF5, 0xF6, 0xF8) if r % 2 else RGBColor(0xFF, 0xFF, 0xFF)
    return tbl


# --------------------------------------------------------------------------- #
def build(out_path: str) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    assets = os.path.join(ROOT, "docs", "ppt", "assets")

    # 竖排三帧的长条图（900x2114）在 PPT 里会被压成细柱 —— 拆出第一帧用
    v1_frame0 = crop_region(
        os.path.join(FIG, "01_dataset_v1/v1_fullframe_gt_projection.jpg"),
        os.path.join(assets, "v1_fullframe_frame0.jpg"), (0.0, 0.0, 1.0, 0.325))

    # ---- 1 封面 ----
    s = blank(prs)
    bar(s, h=Inches(2.35))
    tb = s.shapes.add_textbox(Inches(0.9), Inches(0.75), SLIDE_W - Inches(1.8), Inches(1.6))
    tf = tb.text_frame
    tf.word_wrap = True
    _font(tf.paragraphs[0].add_run(), 38, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    tf.paragraphs[0].runs[0].text = "单目单图 6D 物体位姿估计"
    p = tf.add_paragraph()
    _font(p.add_run(), 20, color=RGBColor(0xD8, 0xE4, 0xFF))
    p.runs[0].text = "目标物体：DJI Osmo Action 4　|　交付物：8 个角点的 2D 投影"
    bullets(s, [
        ("从零自建：三维模型生成 → 合成数据渲染 → 网络实现 → 训练 → 真实视频测试", MUTED, False),
        ("自建数据集上可用（4.12 px / PCK@0.05 96.9%）；目标视频上尚不可用", INK, True),
    ], Inches(0.9), Inches(2.75), Inches(7.2), Inches(1.6), size=17, space=10)
    picture_fit(s, os.path.join(ROOT, "photo_of_the_project.png"),
                (Inches(8.6), Inches(2.65), Inches(4.2), Inches(4.4)))

    # ---- 2 任务与交付物 ----
    s = blank(prs); bar(s)
    title(s, "任务与交付物", "先对齐口径：交付的是 2D 角点，不是 6D 位姿数字")
    bullets(s, [
        ("输入：单张 RGB 图像（目标视频 3248×2464）", INK, False),
        ("输出：目标物体 3D 包围盒 8 个顶点在图像上的 2D 投影", ACCENT, True),
        ("⇒ 网络与损失都不消费相机内参 K —— 本项目不需要标定", GOOD, True),
        ("为什么用角点做中间表示：直接回归旋转量在图像平面上高度非线性；"
         "得到 2D–3D 对应后可用 solvePnP 解位姿，数值上稳定得多", INK, False),
        ("⚠️ 因此评价指标用 2D 角点口径（像素误差 / PCK / 框 IoU），"
         "不用 ADD / ADD-S / 5cm5° —— 那些需要内参跑 PnP，会把“角点其实还行”读成“完全失败”", WARN, False),
    ], Inches(0.7), Inches(1.65), Inches(11.9), Inches(4.6), size=16, space=14)

    # ---- 3 流水线 ----
    s = blank(prs); bar(s)
    title(s, "五步流水线", "① 三维模型　② 渲染合成数据　③ 角点网络　④ 训练　⑤ 目标视频测试")
    bullets(s, [
        ("① 4 张官方图 → 三维生成 → BOP 格式模型", INK, False),
        ("　（三轴 99%/102%/101% 对官方尺寸）", MUTED, False),
        ("② BlenderProc 渲染 → RGB + GT 6D 位姿", INK, False),
        ("③ ResNet18 + FPN 热图解码头，12.0 M", INK, False),
        ("　256×256 → 8×64×64", MUTED, False),
        ("④ 纯合成数据训练（RTX 4090，6.9 步/秒）", INK, False),
        ("⑤ 目标视频测试：GroundingDINO 出框", INK, False),
        ("　→ 网络出 8 角点", MUTED, False),
    ], Inches(0.7), Inches(1.55), Inches(6.2), Inches(3.3), size=14, space=6)
    picture_fit(s, os.path.join(ROOT, "photo_of_the_project.png"),
                (Inches(7.2), Inches(1.50), Inches(5.6), Inches(3.5)))
    picture_fit(s, os.path.join(RES, "stage1_generation/final_4view_hybrid/hy_bop_sheet.png"),
                (Inches(0.7), Inches(4.95), Inches(11.9), Inches(2.05)),
                caption="阶段①产出的 DJI Action 4 三维模型（BOP 渲染，多角度 × 多场景）")

    # ---- 4 数据集主线（1/2）----
    s = blank(prs); bar(s)
    title(s, "数据集主线（1/2）：第一版", "多个 DJI 随机摆放在盒子里 —— 想到用“同类堆叠”制造遮挡")
    picture_fit(s, v1_frame0,
                (Inches(0.7), Inches(1.65), Inches(6.8), Inches(5.2)),
                caption="绿框 = 目标 bbox_visib；红圈 = 目标 8 角点投影；灰框 = 其他同类实例")
    bullets(s, [
        ("规模：25 场景 × 20 帧 = 500 帧 / 5000 实例", INK, True),
        ("GT：物体坐标系下 3D 盒的 8 角点，按 GT 位姿投影到图像", INK, False),
        ("数据集本身没问题（热图峰与角点 0 格误差）", GOOD, False),
        ("", INK, False),
        ("⚠️ 但埋了一个隐患", WARN, True),
        ("画面里每一个物体都是同一款 DJI", WARN, False),
        ("（图中灰框全是 DJI）", WARN, False),
        ("⇒ 模型无法学“哪个是 DJI”，", WARN, False),
        ("　只能学“框里居中的那个”", WARN, False),
        ("⇒ 1.4× 裁剪框内 96.2% 含其他同类实例", WARN, False),
    ], Inches(7.8), Inches(1.70), Inches(4.8), Inches(5.1), size=14, space=8)

    # ---- 5 数据集主线（2/2）----
    s = blank(prs); bar(s)
    title(s, "数据集主线（2/2）：第二版", "针对“渲染得太干净”改造 —— 屏幕有画面 + 加异类实例")
    picture_fit(s, os.path.join(FIG, "02_dataset_geo7/geo7_scene000000_f0000.png"),
                (Inches(0.7), Inches(1.55), Inches(5.9), Inches(3.55)),
                caption="屏幕内容随机：一台显示画面，其余显示菜单界面")
    picture_fit(s, os.path.join(FIG, "02_dataset_geo7/geo_series_static_cover_with_gt.jpg"),
                (Inches(6.75), Inches(1.55), Inches(5.9), Inches(3.55)),
                caption="带 GT 3D 框投影；Objaverse 异类道具压在 DJI 上造遮挡")
    bullets(s, [
        ("三处改动：① DJI 屏幕上贴一个平面 → 模拟“正在拍摄”（屏幕有画面）"
         "　② 加入从 Objaverse 爬下的其它实例 → 外观多样性 + 自然遮挡"
         "　③ 随机资产 / 随机屏幕内容 / 随机色温曝光 → 逼模型学几何而非外观", INK, False),
        ("规模：700 帧 / 4200 实例", INK, True),
        ("代价：合成集分数会掉（见“排名反转”），但目标视频上反而更好", WARN, False),
    ], Inches(0.7), Inches(5.45), Inches(11.9), Inches(1.7), size=14, space=8)

    # ---- 6 模型主线 ----
    s = blank(prs); bar(s)
    title(s, "模型主线", "架构：ResNet18 + FPN 热图解码头 · 12.0 M 参数 · 256×256 → 8×64×64")
    table(s, [
        ["版本", "改动", "结果"],
        ["v1", "RGB → ResNet → 热图 → 角点（基线）", "自建数据集效果好；目标视频一般"],
        ["E6 / E7", "把 ResNet 换成 DINOv2 / 改输入尺寸", "❌ 没用：DINOv2 反而更差（23.3 vs 14.4 px）"],
        ["GEO7", "沿用同架构，改装好的第二版数据", "✅ 目标视频优于 v1（IoU 0.799 vs 0.500），但仍不够"],
        ["MIGEO7", "多实例：中心热图 + 角点偏移", "合成集 recall 0.83；真实帧上不可靠"],
        ["GEO8", "加几何约束 / 增强", "❌ 增强是负收益（同 epoch 落后 10~25%）"],
    ], Inches(0.7), Inches(1.60), Inches(11.9), Inches(3.1), col_w=[1.0, 4.2, 6.0], size=13)
    bullets(s, [
        ("最关键的一条教训：", WARN, True),
        ("前几版都在追“自建数据集上的分数”，但真正的目标是目标视频 —— 优化目标选错了，"
         "换 DINOv2 / 调参数自然都没用。", INK, False),
        ("后经老师提示，确认根因是「训练数据与真实场景差别过大」。", INK, False),
    ], Inches(0.7), Inches(5.00), Inches(11.9), Inches(2.0), size=15, space=9)

    # ---- 7 效果（1/2）自建数据集 ----
    s = blank(prs); bar(s)
    title(s, "效果（1/2）：自建数据集上 —— 可用", "独立 fp32 评测，248 样本 / 1984 角点")
    picture_fit(s, os.path.join(FIG, "03_v1_on_dataset/v1_best4_worst4.png"),
                (Inches(0.7), Inches(1.65), Inches(8.3), Inches(5.0)),
                caption="上排 BEST 4（1.0~1.4 px，visb=1.00）／下排 WORST 4（71.7~141.9 px，visb≈0.25）")
    bullets(s, [
        ("mean 4.12 px", GOOD, True),
        ("median 1.44 px", INK, False),
        ("PCK@0.05　96.88%", GOOD, True),
        ("框 IoU　0.961", INK, False),
        ("", INK, False),
        ("⚠️ 失败集中在严重遮挡", WARN, True),
        ("完全可见样本失败率 0.79%，", INK, False),
        ("严重遮挡样本 45.0%", WARN, True),
        ("差距 57 倍 —— 训练集 90% 完全可见，模型没学过外推", INK, False),
    ], Inches(9.3), Inches(1.75), Inches(3.4), Inches(5.0), size=14, space=8)

    # ---- 8 效果（2/2）目标视频 ----
    s = blank(prs); bar(s)
    title(s, "效果（2/2）：目标视频上 —— 还不可用", "同一批帧，三个模型并排（只留裁剪放大）")
    picture_fit(s, os.path.join(FIG, "06_geo7_on_video/video_3model_zoom_compare.png"),
                (Inches(0.7), Inches(1.55), Inches(3.8), Inches(5.3)),
                caption="列 = train_v1 / GEO1 / GEO7\n行 = 好样本#1 / #2 / 失败样本")
    table(s, [
        ["模型", "自建数据集", "目标视频 IoU"],
        ["train_v1", "4.12 px（最好）", "0.500（最差）"],
        ["GEO1", "16.37 px", "0.605"],
        ["GEO7", "39.21 px（最差）", "0.799（最好）"],
    ], Inches(4.8), Inches(1.65), Inches(7.8), Inches(1.7), col_w=[1.4, 2.9, 2.9], size=13)
    bullets(s, [
        ("⚠️ 排名完全反转", WARN, True),
        ("自建数据集上最烂的 GEO7（差 9.5 倍），在目标视频上最好", INK, False),
        ("⇒ 自建数据集的绝对分数不能作为“能不能用”的判据", WARN, True),
        ("", INK, False),
        ("原因：v1 的测试集与训练集共享同一套外观统计 → 分数虚高；"
         "第二版用随机资产/屏幕/色温，逼模型学几何 → 合成集掉分但真实域更准", INK, False),
        ("", INK, False),
        ("⚠️ 目标视频目前只有 1~2 帧手工标注 —— 是强信号，不是定论", MUTED, False),
    ], Inches(4.8), Inches(3.60), Inches(7.8), Inches(3.2), size=13.5, space=7)

    # ---- 9 已知缺陷 ----
    s = blank(prs); bar(s)
    title(s, "已知缺陷与根因", "都定位到了具体机制，不是“效果不好”这种模糊描述")
    bullets(s, [
        ("① 域差距（最主要）：真实的目标 DJI 屏幕点亮、机身贴有 QR 标签，"
         "而我们渲染的是干净模型 → 外观对不上", WARN, True),
        ("② 预测的 8 个角点不构成长方体：四条深度棱夹角 72°~90°（合法投影应 <10°）。"
         "根因是架构 —— 8 张热图各取各的最大值，损失也是逐角点的，"
         "没有任何一项在问“这 8 点合起来像不像盒子”", WARN, True),
        ("③ 数据设计：每帧每个物体都是同一款 DJI（v1 每帧 10 个、第二版 6 个），"
         "裁剪框 96.2% 含其他同类实例", INK, False),
        ("④ 训练集遮挡不足：90% 实例完全可见 → 严重遮挡处失败率高 57 倍", INK, False),
        ("⑤ 目标视频没有 GT 角点标注 → “有没有变好”无法量化", INK, False),
        ("⑥ 多实例路线：合成集 recall 0.83，但真实帧上框的松紧与位置都不可靠", INK, False),
    ], Inches(0.7), Inches(1.60), Inches(11.9), Inches(5.0), size=15, space=13)

    # ---- 10 教训与下一步 ----
    s = blank(prs); bar(s)
    title(s, "教训与下一步")
    bullets(s, [
        ("教训（自己总结的）", INK, True),
        ("· 做好实验记录；开始跑实验之前充分讨论，否则浪费 GPU 资源", ACCENT, True),
        ("· 优化目标要选对：真正的目标是目标视频，不是自建数据集上的分数", ACCENT, True),
        ("· 以为“改网络结构/调参数”能解决，实则数据的问题占了主导", INK, False),
        ("", INK, False),
        ("下一步（按性价比排序）", INK, True),
        ("1. 渲染时对齐真实外观：屏幕点亮 + 给目标贴标签 —— 唯一直接针对域差距的动作", INK, False),
        ("2. 同一批渲染里改数据设计：每帧 1 个 DJI + 10 个异类道具", INK, False),
        ("3. 给目标视频建 GT（哪怕 8~10 帧），否则无法量化改进", INK, False),
        ("4. 给角点加几何约束（位姿参数化，或 PnP 后处理重投影）", INK, False),
        ("5. 用 BoxDreamer 出位姿再重投影我们的 CAD 盒 —— 比自训便宜，值得先试", INK, False),
    ], Inches(0.7), Inches(1.55), Inches(11.9), Inches(5.4), size=14.5, space=8)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    prs.save(out_path)
    print(f"已写出 {out_path}  ({len(prs.slides.__iter__.__self__._sldIdLst)} 页, "
          f"{os.path.getsize(out_path)/1e6:.2f} MB)")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(ROOT, "docs", "ppt", "minimal.pptx"))
    args = ap.parse_args()
    build(args.out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
